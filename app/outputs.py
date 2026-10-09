"""Fișiere de ieșire suplimentare: DOCX bilingv (pentru verificare) și raportul de calitate (HTML)."""

from __future__ import annotations

import datetime as _dt
import html
from pathlib import Path

from .pdf_extract import PageInfo


def write_bilingual_docx(pages: list[PageInfo], path: Path, meta: dict) -> Path:
    """Tabel: Nr | Pagină | Text sursă | Traducere RO | Observații (toate segmentele)."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Cm

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.5))

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(9)

    doc.add_heading(f"Traducere bilingvă - {meta.get('file', '')}", level=1)
    doc.add_paragraph(
        f"Motor: {meta.get('engine', '')} | Limba sursă: {meta.get('lang', 'auto')} | "
        f"Data: {meta.get('date', '')} | Segmente: {meta.get('total', 0)} | "
        f"Avertismente: {meta.get('warnings', 0)}"
    )

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    widths = [Cm(1.2), Cm(1.4), Cm(10.5), Cm(10.5), Cm(3.5)]
    hdr = table.rows[0].cells
    for cell, text, w in zip(hdr, ["Nr.", "Pag.", "Text sursă", "Traducere în română", "Observații"], widths):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        cell.width = w
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "D9E2F3")
        cell._tc.get_or_add_tcPr().append(shading)

    for pinfo in pages:
        for seg in pinfo.segments:
            row = table.add_row().cells
            values = [
                str(seg.id + 1), str(pinfo.index + 1), seg.text,
                seg.translation if not seg.skipped else "(nemodificat)",
                "; ".join(seg.warnings),
            ]
            for cell, val, w in zip(row, values, widths):
                cell.text = val
                cell.width = w
                if seg.warnings:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), "FFF2CC")
                    cell._tc.get_or_add_tcPr().append(shading)
    doc.save(str(path))
    return path


def write_report_html(pages: list[PageInfo], path: Path, meta: dict) -> Path:
    """Raport de calitate: rezumat + segmentele cu avertismente (numere, marcaje, glosar, micșorare)."""
    rows = []
    n_warn = 0
    for pinfo in pages:
        for seg in pinfo.segments:
            if seg.warnings:
                n_warn += 1
                rows.append(
                    "<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td class='w'>{}</td></tr>".format(
                        seg.id + 1, pinfo.index + 1, html.escape(seg.text), html.escape(seg.translation),
                        "<br>".join(html.escape(w) for w in seg.warnings),
                    )
                )
    scanned = [p.index + 1 for p in pages if p.is_scanned]
    stats = meta.get("stats", {})
    body = f"""<!doctype html><html lang="ro"><head><meta charset="utf-8">
<title>Raport traducere - {html.escape(meta.get('file', ''))}</title>
<style>
 body {{ font-family: Segoe UI, Arial, sans-serif; margin: 24px; color: #222; }}
 h1 {{ font-size: 20px; }} h2 {{ font-size: 16px; margin-top: 28px; }}
 table {{ border-collapse: collapse; width: 100%; font-size: 13px; }}
 th, td {{ border: 1px solid #ccc; padding: 6px 8px; vertical-align: top; }}
 th {{ background: #d9e2f3; text-align: left; }}
 td.w {{ background: #fff2cc; }}
 .ok {{ color: #1a7f37; font-weight: bold; }} .bad {{ color: #b3261e; font-weight: bold; }}
 .meta td {{ border: none; padding: 2px 10px 2px 0; }}
</style></head><body>
<h1>Raport de traducere - {html.escape(meta.get('file', ''))}</h1>
<table class="meta">
<tr><td>Data</td><td>{html.escape(meta.get('date', ''))}</td></tr>
<tr><td>Motor</td><td>{html.escape(meta.get('engine', ''))}{' + post-editare Claude' if meta.get('post_edit') else ''}</td></tr>
<tr><td>Limba sursă</td><td>{html.escape(meta.get('lang', 'auto'))}</td></tr>
<tr><td>Pagini</td><td>{len(pages)}</td></tr>
<tr><td>Segmente</td><td>{meta.get('total', 0)} (traduse: {meta.get('translated', 0)}, nemodificate: {meta.get('kept', 0)}, din memorie: {meta.get('from_tm', 0)})</td></tr>
<tr><td>Reinserate în PDF</td><td>{stats.get('replaced', 0)} (micșorate sub prag: {stats.get('shrunk', 0)}, erori: {stats.get('errors', 0)})</td></tr>
<tr><td>Segmente cu avertismente</td><td class="{'ok' if n_warn == 0 else 'bad'}">{n_warn}</td></tr>
</table>
"""
    if scanned:
        body += ("<p class='bad'>Pagini fără strat de text (probabil scanate), NETRADUSE: "
                 + ", ".join(map(str, scanned)) + ". Aplicați OCR înainte de traducere.</p>")
    body += "<h2>Segmente de verificat</h2>"
    if rows:
        body += ("<p>Verificați manual segmentele de mai jos: numerele sau codurile diferă, un marcaj "
                 "protejat s-a pierdut, un termen din glosar nu a fost aplicat sau textul a fost micșorat mult.</p>"
                 "<table><tr><th>Nr.</th><th>Pag.</th><th>Sursă</th><th>Traducere</th><th>Avertisment</th></tr>"
                 + "".join(rows) + "</table>")
    else:
        body += "<p class='ok'>Nicio problemă detectată automat.</p>"
    body += "</body></html>"
    path.write_text(body, encoding="utf-8")
    return path


def now_str() -> str:
    return _dt.datetime.now().strftime("%Y-%m-%d %H:%M")
